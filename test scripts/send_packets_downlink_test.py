import threading
import time
from datetime import datetime
from pathlib import Path

import grpc
from chirpstack_api import api
import paho.mqtt.client as mqtt
from openpyxl import Workbook

# ---------------- CONFIGURATION ----------------
# Replace these placeholders with your specific setup details
SERVER_IP = "YOUR_SERVER_IP:8080"
MQTT_BROKER = "YOUR_MQTT_BROKER_IP"
MQTT_PORT = 1883

# Device and Security credentials
DEV_EUI = "YOUR_DEVICE_EUI"
API_TOKEN = "YOUR_API_TOKEN"

# MQTT Topic for monitoring downlink commands
MQTT_TOPIC = "eu868/gateway/+/command/down"

# Test Parameters
NUM_PACKETS = 20

# Output Path (Saves to Desktop)
DESKTOP = Path.home() / "Desktop"
LOG_FILE = DESKTOP / "downlink_gateway_log.xlsx"
# -----------------------------------------------

gateway_counter = {}
records = []
received_packets = 0
stop_event = threading.Event()

def on_connect(client, userdata, flags, rc):
    print("Connected to MQTT broker")
    client.subscribe(MQTT_TOP_IC)
    print("Subscribed to:", MQTT_TOPIC)

def on_message(client, userdata, msg):
    global received_packets
    global gateway_counter
    global records

    try:
        topic_parts = msg.topic.split("/")
        gateway_id = topic_parts[2]
        timestamp = datetime.now().isoformat()

        received_packets += 1

        if gateway_id not in gateway_counter:
            gateway_counter[gateway_id] = 0

        gateway_counter[gateway_id] += 1
        records.append([timestamp, gateway_id])

        print(f"\nDownlink {received_packets}/{NUM_PACKETS}")
        print("Gateway selected:", gateway_id)

        if received_packets >= NUM_PACKETS:
            print("\nAll packets received. Stopping script.")
            stop_event.set()

    except Exception as e:
        print("Error processing message:", e)

def start_mqtt():
    client = mqtt.Client()
    client.on_connect = on_connect
    client.on_message = on_message
    client.connect(MQTT_BROKER, MQTT_PORT, 60)
    client.loop_start()
    return client

def send_downlinks():
    channel = grpc.insecure_channel(SERVER_IP)
    client = api.DeviceServiceStub(channel)
    auth_token = [("authorization", "Bearer %s" % API_TOKEN)]

    for i in range(NUM_PACKETS):
        payload = bytes([i + 1])
        req = api.EnqueueDeviceQueueItemRequest()
        req.queue_item.confirmed = False
        req.queue_item.data = payload
        req.queue_item.dev_eui = DEV_EUI
        req.queue_item.f_port = 10

        resp = client.Enqueue(req, metadata=auth_token)
        print(f"\nDownlink {i+1}/{NUM_PACKETS} enqueued")
        print("Downlink ID:", resp.id)
        time.sleep(2)

def save_excel():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Downlink Log"
    ws1.append(["Timestamp", "Gateway"])
    for r in records:
        ws1.append(r)

    ws2 = wb.create_sheet("Statistics")
    ws2.append(["Gateway", "Packets", "Percentage"])
    total = sum(gateway_counter.values())

    for gw, count in gateway_counter.items():
        percentage = (count / total) * 100
        ws2.append([gw, count, percentage])

    wb.save(LOG_FILE)
    print(f"\nExcel file saved at:\n{LOG_FILE}")

def print_gateway_statistics():
    print("\n------ Gateway Usage Statistics ------")
    total = sum(gateway_counter.values())
    for gw, count in gateway_counter.items():
        percentage = (count / total) * 100
        print(f"{gw} -> {count} packets ({percentage:.2f}%)")

if __name__ == "__main__":
    mqtt_client = start_mqtt()
    time.sleep(2)
    send_downlinks()
    print("\nWaiting for gateway transmissions...")
    stop_event.wait()
    mqtt_client.loop_stop()
    save_excel()
    print_gateway_statistics()
    print("Script finished.")
